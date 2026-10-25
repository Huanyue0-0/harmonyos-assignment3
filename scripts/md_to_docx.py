#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把作业报告 Markdown 转换为同名 .docx。

用法：
    python scripts/md_to_docx.py docs/report/Assignment3_报告.md docs/report/Assignment3_报告.docx

依赖：python-docx（若系统 Python 受 PEP 668 保护，请使用虚拟环境安装）。
支持的 Markdown 子集：# ~ #### 标题、段落、- 无序列表、1. 有序列表、
> 引用、```代码块```、| 表格 |、**加粗**、`行内代码`、--- 分隔线。
"""

import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def set_default_font(doc: Document) -> None:
    """设置中文字体，避免 Word 打开后中文回退成方框。"""
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')


def add_inline(paragraph, text: str) -> None:
    """处理 **加粗** 与 `行内代码`。"""
    pattern = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`)')
    for piece in pattern.split(text):
        if not piece:
            continue
        if piece.startswith('**') and piece.endswith('**') and len(piece) > 4:
            run = paragraph.add_run(piece[2:-2])
            run.bold = True
        elif piece.startswith('`') and piece.endswith('`') and len(piece) > 2:
            run = paragraph.add_run(piece[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        else:
            paragraph.add_run(piece)


def split_table_row(line: str) -> list:
    """按未转义的 | 切分表格行，并还原 \\| 。"""
    cells = []
    buf = []
    escaped = False
    for ch in line.strip().strip('|'):
        if escaped:
            buf.append(ch)
            escaped = False
            continue
        if ch == '\\':
            escaped = True
            continue
        if ch == '|':
            cells.append(''.join(buf).strip())
            buf = []
            continue
        buf.append(ch)
    cells.append(''.join(buf).strip())
    return cells


def is_separator_row(line: str) -> bool:
    stripped = line.strip().strip('|')
    return bool(stripped) and set(stripped.replace('|', '').replace(' ', '')) <= set('-:')


def add_table(doc: Document, rows: list) -> None:
    header = split_table_row(rows[0])
    body = [split_table_row(r) for r in rows[1:] if not is_separator_row(r)]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, text in enumerate(header):
        cell = table.rows[0].cells[index]
        cell.text = ''
        add_inline(cell.paragraphs[0], text)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in body:
        cells = table.add_row().cells
        for index in range(len(header)):
            value = row[index] if index < len(row) else ''
            cells[index].text = ''
            add_inline(cells[index].paragraphs[0], value)
    doc.add_paragraph()


def convert(md_path: str, docx_path: str) -> None:
    with open(md_path, 'r', encoding='utf-8') as handle:
        lines = handle.read().split('\n')

    doc = Document()
    set_default_font(doc)

    index = 0
    total = len(lines)
    while index < total:
        line = lines[index]
        stripped = line.rstrip()

        # 代码块
        if stripped.startswith('```'):
            index += 1
            buffer = []
            while index < total and not lines[index].rstrip().startswith('```'):
                buffer.append(lines[index])
                index += 1
            index += 1
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(12)
            run = paragraph.add_run('\n'.join(buffer))
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            continue

        # 标题
        if stripped.startswith('#'):
            level = len(stripped) - len(stripped.lstrip('#'))
            text = stripped[level:].strip().replace('`', '').replace('**', '')
            doc.add_heading(text, level=min(level, 4))
            index += 1
            continue

        # 分隔线
        if stripped.strip() == '---':
            doc.add_paragraph()
            index += 1
            continue

        # 表格
        if stripped.startswith('|') and index + 1 < total and is_separator_row(lines[index + 1]):
            rows = []
            while index < total and lines[index].strip().startswith('|'):
                rows.append(lines[index])
                index += 1
            add_table(doc, rows)
            continue

        # 引用
        if stripped.startswith('>'):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(18)
            add_inline(paragraph, stripped.lstrip('>').strip())
            for run in paragraph.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            index += 1
            continue

        # 无序列表
        if re.match(r'^\s*[-*]\s+', line):
            text = re.sub(r'^\s*[-*]\s+', '', line)
            paragraph = doc.add_paragraph(style='List Bullet')
            add_inline(paragraph, text)
            index += 1
            continue

        # 有序列表
        if re.match(r'^\s*\d+\.\s+', line):
            text = re.sub(r'^\s*\d+\.\s+', '', line)
            paragraph = doc.add_paragraph(style='List Number')
            add_inline(paragraph, text)
            index += 1
            continue

        # 空行
        if stripped.strip() == '':
            index += 1
            continue

        # 普通段落
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        add_inline(paragraph, stripped.strip())
        index += 1

    doc.save(docx_path)
    print(f'生成成功: {docx_path}')


if __name__ == '__main__':
    if len(sys.argv) != 3:
        print('用法: python md_to_docx.py <input.md> <output.docx>')
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2])
